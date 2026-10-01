# -*- coding: utf-8 -*-
"""排版体检：检查越界与文本溢出"""
import math
from pptx import Presentation
from pptx.util import Emu

EMU_IN = 914400
PW, PH = 13.333, 7.5


def char_w(ch, size):
    """估算字符宽度(pt)"""
    if ord(ch) > 0x2E80:          # CJK 及全角
        return size * 1.0
    if ch == " ":
        return size * 0.30
    return size * 0.55


def est_lines(text, width_in, size, spacing=1.0):
    wpt = width_in * 72.0
    rows = 0
    for seg in text.split("\n"):
        if not seg:
            rows += 1
            continue
        total = sum(char_w(c, size) for c in seg)
        rows += max(1, math.ceil(total / max(wpt, 1)))
    return rows


def check(path):
    prs = Presentation(path)
    prs_w = prs.slide_width / EMU_IN
    prs_h = prs.slide_height / EMU_IN
    problems = []
    for si, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            L, T = sh.left / EMU_IN, sh.top / EMU_IN
            Wd, Ht = sh.width / EMU_IN, sh.height / EMU_IN
            # 1. 越界
            if L < -0.02 or T < -0.02 or L + Wd > prs_w + 0.02 or T + Ht > prs_h + 0.02:
                problems.append(f"  P{si:02d} 越界  {sh.shape_type}  L={L:.2f} T={T:.2f} "
                                f"R={L+Wd:.2f} B={T+Ht:.2f}")
            # 2. 文本溢出
            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            need = 0.0
            size = 12.0
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if not txt:
                    continue
                sz = max((r.font.size.pt for r in p.runs if r.font.size), default=12.0)
                size = sz
                sp = p.line_spacing or 1.0
                if isinstance(sp, float) is False:
                    sp = 1.0
                ml = (tf.margin_left or 0) / EMU_IN
                mr = (tf.margin_right or 0) / EMU_IN
                avail = Wd - ml - mr
                if avail <= 0.1:
                    continue
                n = est_lines(txt, avail, sz)
                need += n * sz * sp * 1.22 / 72.0
            mt = (tf.margin_top or 0) / EMU_IN
            mb = (tf.margin_bottom or 0) / EMU_IN
            if need > 0 and need > Ht - mt - mb + 0.06:
                body = "".join(r.text for p in tf.paragraphs for r in p.runs)[:34]
                problems.append(f"  P{si:02d} 溢出 需{need+mt+mb:.2f}in / 框{Ht:.2f}in "
                                f"[{size:.0f}pt] 「{body}…」")
    return problems


if __name__ == "__main__":
    import sys
    p = sys.argv[1] if len(sys.argv) > 1 else "五千万美元的深夜_十一家AI决策实验.pptx"
    probs = check(p)
    if not probs:
        print("PASS 无越界、无溢出")
    else:
        print(f"发现 {len(probs)} 处：")
        for x in probs:
            print(x)
